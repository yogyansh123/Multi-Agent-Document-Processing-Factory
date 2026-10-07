"""
services/ocr/tesseract.py
=========================
Concrete OCR provider using Tesseract, PyMuPDF, and python-docx.

Features:
- DOCX: Native text extraction using python-docx (fast, no OCR overhead).
- PDF: Two-phase extraction:
    1. Fast path: Extract embedded digital text via PyMuPDF.
    2. Fallback: If page has no embedded text (scanned PDF), render page
       to image at configured DPI and run pytesseract.
- Images (PNG, JPG, JPEG, TIFF, BMP, WEBP): Direct pytesseract extraction.
- Fully asynchronous: All CPU/disk bound operations are offloaded to
  worker threads using asyncio.to_thread.
- Clean error mapping: pytesseract.TesseractNotFoundError and other exceptions
  are mapped to OCRError with clear, actionable messages.
"""

from __future__ import annotations

import asyncio
import time
from pathlib import Path
from typing import Any

from PIL import Image
import pytesseract

from app.core.config import settings
from app.services.ocr.base import OCRError, OCRProvider, OCRResult


class TesseractProvider(OCRProvider):
    """
    OCR provider leveraging local Tesseract OCR engine and PyMuPDF.
    """

    def __init__(
        self,
        tesseract_cmd: str | None = None,
        language: str | None = None,
        dpi: int | None = None,
    ) -> None:
        self._tesseract_cmd = tesseract_cmd or settings.TESSERACT_CMD
        self._language = language or settings.OCR_LANGUAGE
        self._dpi = dpi or settings.OCR_DPI

        if self._tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = self._tesseract_cmd

    @property
    def provider_name(self) -> str:
        return "tesseract"

    async def extract_text(self, file_path: str, file_type: str) -> OCRResult:
        """
        Extract text from file asynchronously.
        """
        path = Path(file_path)
        if not path.is_file():
            raise OCRError(f"File not found: {file_path}", provider=self.provider_name)

        ext = file_type.lower().lstrip(".")

        try:
            return await asyncio.to_thread(self._process_file, str(path), ext)
        except OCRError:
            raise
        except pytesseract.TesseractNotFoundError as exc:
            raise OCRError(
                "Tesseract executable not found. Please install Tesseract or configure TESSERACT_CMD.",
                provider=self.provider_name,
                cause=exc,
            ) from exc
        except Exception as exc:
            raise OCRError(
                f"OCR processing failed for {path.name}: {exc}",
                provider=self.provider_name,
                cause=exc,
            ) from exc

    def _process_file(self, file_path: str, ext: str) -> OCRResult:
        start_time = time.monotonic()

        if ext == "docx":
            result = self._process_docx(file_path)
        elif ext == "pdf":
            result = self._process_pdf(file_path)
        elif ext in {"png", "jpg", "jpeg", "tiff", "bmp", "webp"}:
            result = self._process_image(file_path)
        else:
            raise OCRError(
                f"Unsupported file type for OCR: '{ext}'",
                provider=self.provider_name,
            )

        elapsed_ms = max(1, int((time.monotonic() - start_time) * 1000))
        return OCRResult(
            text=result["text"],
            provider=self.provider_name,
            page_count=result["page_count"],
            processing_time_ms=elapsed_ms,
            metadata=result.get("metadata", {}),
        )

    def _process_docx(self, file_path: str) -> dict[str, Any]:
        import docx

        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        table_lines = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_lines.append(row_text)

        full_text = "\n".join(paragraphs + table_lines).strip()
        return {
            "text": full_text,
            "page_count": 1,
            "metadata": {
                "format": "docx",
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "extracted_via": "native_docx",
            },
        }

    def _process_image(self, file_path: str) -> dict[str, Any]:
        with Image.open(file_path) as img:
            text = pytesseract.image_to_string(img, lang=self._language)
            return {
                "text": text.strip(),
                "page_count": 1,
                "metadata": {
                    "format": "image",
                    "language": self._language,
                    "image_size": list(img.size),
                },
            }

    def _process_pdf(self, file_path: str) -> dict[str, Any]:
        import fitz  # PyMuPDF

        doc = fitz.open(file_path)
        page_count = len(doc)
        pages_text: list[str] = []
        pages_scanned_count = 0
        pages_digital_count = 0

        zoom = self._dpi / 72.0
        matrix = fitz.Matrix(zoom, zoom)

        for page_idx in range(page_count):
            page = doc[page_idx]
            page_text = page.get_text()

            # If page contains reasonable text, use digital extraction
            if page_text and len(page_text.strip()) >= 10:
                pages_text.append(page_text.strip())
                pages_digital_count += 1
            else:
                # Render to pixmap and run pytesseract
                pix = page.get_pixmap(matrix=matrix)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                ocr_text = pytesseract.image_to_string(img, lang=self._language)
                pages_text.append(ocr_text.strip())
                pages_scanned_count += 1

        doc.close()

        full_text = "\n\n".join(pages_text).strip()
        return {
            "text": full_text,
            "page_count": page_count,
            "metadata": {
                "format": "pdf",
                "language": self._language,
                "dpi": self._dpi,
                "digital_pages": pages_digital_count,
                "scanned_pages": pages_scanned_count,
            },
        }
